import hashlib
import io
import math
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from uuid import uuid4
import pandas as pd
from openpyxl.utils import get_column_letter
from .models import Mapping
from . import store

FIELDS = {
    'inventory': ['sku', 'item', 'supplier', 'on_hand', 'lead_days', 'unit', 'unit_cost', 'currency'],
    'sales': ['sku', 'date', 'qty', 'unit'],
    'orders': ['order_id', 'sku', 'supplier', 'expected_date', 'ordered_qty', 'received_qty', 'unit'],
    'invoices': ['invoice_id', 'supplier', 'date', 'amount', 'currency'],
    'stock_history': ['sku', 'date', 'on_hand', 'lead_days', 'supplier', 'unit'],
}
ALIASES = {
    'sku': ['sku', 'itemcode', 'productid', 'productcode'], 'item': ['item', 'product', 'itemname', 'productname', 'description'],
    'supplier': ['supplier', 'vendor', 'suppliername', 'vendorname'], 'on_hand': ['onhand', 'stock', 'quantityonhand', 'stockqty'],
    'lead_days': ['leaddays', 'leadtime', 'leadtimeindays'], 'unit': ['unit', 'uom', 'units'],
    'unit_cost': ['unitcost', 'cost', 'costprice'], 'currency': ['currency', 'ccy'], 'date': ['date', 'saledate', 'invoicedate'],
    'qty': ['qty', 'quantity', 'sold', 'quantitysold'], 'order_id': ['orderid', 'po', 'ponumber', 'purchaseorder'],
    'expected_date': ['expecteddate', 'duedate', 'expecteddelivery'], 'ordered_qty': ['orderedqty', 'ordered', 'quantityordered'],
    'received_qty': ['receivedqty', 'received', 'quantityreceived'], 'invoice_id': ['invoiceid', 'invoicenumber', 'invoice'],
    'amount': ['amount', 'total', 'invoicetotal', 'expense'],
}

def normalize(value):
    return re.sub(r'[^a-z0-9]', '', str(value).lower())

def suggest(headers):
    result = {}
    for field, aliases in ALIASES.items():
        match = next((h for h in headers if normalize(h) in aliases), None)
        if match:
            result[field] = match
    scores = {kind: len(set(fields) & set(result)) / len(fields) for kind, fields in FIELDS.items()}
    kind = max(scores, key=scores.get)
    return {'kind': kind, 'columns': {f: result[f] for f in FIELDS[kind] if f in result}, 'date_format': 'iso', 'currency': 'USD'}

def parse_file(content: bytes, name: str, synthetic=False):
    if len(content) > 15 * 1024 * 1024:
        raise ValueError('Maximum upload size is 15 MB.')
    suffix = name.lower().rsplit('.', 1)[-1]
    sheet = 'CSV'
    if suffix == 'csv':
        frame = pd.read_csv(io.BytesIO(content), dtype=str, keep_default_na=False, nrows=50001, encoding='utf-8-sig')
    elif suffix == 'xlsx':
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Expanded workbook exceeds 100 MB.')
        book = pd.ExcelFile(io.BytesIO(content), engine='openpyxl')
        if len(book.sheet_names) != 1:
            raise ValueError('Upload one sheet per workbook so evidence has an unambiguous source.')
        sheet = book.sheet_names[0]
        frame = pd.read_excel(book, sheet_name=sheet, dtype=str, keep_default_na=False, nrows=50001)
    else:
        raise ValueError('Use a UTF-8 CSV or an XLSX workbook.')
    if frame.empty or len(frame) > 50000 or len(frame.columns) > 80:
        raise ValueError('Upload 1–50,000 rows and at most 80 columns.')
    headers = [str(h).strip() for h in frame.columns]
    if len(set(headers)) != len(headers) or any(not h or h.startswith('Unnamed:') or re.search(r'\.\d+$', h) for h in headers):
        raise ValueError('Use unique nonblank column headers; remove duplicate header suffixes.')
    frame.columns = headers
    dataset_id = uuid4().hex
    original = store.DATA / 'originals' / dataset_id
    original.parent.mkdir(parents=True, exist_ok=True)
    original.write_bytes(content)
    raw = frame.fillna('').astype(str).to_dict(orient='records')
    result = {'id': dataset_id, 'name': name.replace('\\', '/').split('/')[-1][:150], 'sheet': sheet,
              'sha256': hashlib.sha256(content).hexdigest(), 'headers': headers, 'row_count': len(raw),
              'confirmed': False, 'synthetic': synthetic, 'kind': None, 'mapping': None,
              'suggestion': suggest(headers), 'preview': raw[:5]}
    store.put('raw', {'id': dataset_id, 'rows': raw})
    store.put('datasets', result)
    store.audit('File imported', f"{result['name']}: {len(raw)} rows; original preserved with SHA-256.")
    return result

def parse_date(value, convention):
    text = str(value).strip()
    if convention == 'iso':
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}(?: 00:00:00)?', text):
            raise ValueError('Use ISO YYYY-MM-DD or confirm a different date convention.')
        return date.fromisoformat(text[:10]).isoformat()
    fmt = '%d/%m/%Y' if convention == 'dayfirst' else '%m/%d/%Y'
    return datetime.strptime(text, fmt).date().isoformat()

def validate(dataset, raw, mapping: Mapping):
    fields = FIELDS[mapping.kind]
    if set(mapping.columns) != set(fields):
        raise ValueError('Map every required field, and only fields for the selected dataset type.')
    if len(set(mapping.columns.values())) != len(fields) or any(h not in dataset['headers'] for h in mapping.columns.values()):
        raise ValueError('Each field must map to a different existing column.')
    valid, quality = [], []
    number_fields = {'on_hand', 'lead_days', 'unit_cost', 'qty', 'ordered_qty', 'received_qty', 'amount'}
    seen = {}
    for row_number, source in enumerate(raw, 2):
        row, problems = {}, []
        for field in fields:
            header = mapping.columns[field]
            value = str(source[header]).strip()
            cell = f"{get_column_letter(dataset['headers'].index(header) + 1)}{row_number}"
            try:
                if not value:
                    raise ValueError('Blank required value.')
                if field in number_fields:
                    parsed = float(value)
                    if not math.isfinite(parsed) or parsed < 0 or (field == 'lead_days' and parsed <= 0):
                        raise ValueError('Use a finite nonnegative number; lead time must be greater than zero. Returns/credits require a separate accounting flow.')
                    if parsed > 1_000_000_000_000:
                        raise ValueError('Value exceeds the supported numeric limit of one trillion; check units and formatting.')
                    if field in {'unit_cost', 'amount'} and Decimal(value) != Decimal(value).quantize(Decimal('0.01')):
                        raise ValueError('Monetary values must have at most two decimal places; confirm the accounting precision.')
                    row[field] = parsed
                elif field in {'date', 'expected_date'}:
                    row[field] = parse_date(value, mapping.date_format)
                elif field == 'currency':
                    if value.upper() != mapping.currency:
                        raise ValueError(f'Currency differs from confirmed {mapping.currency}; no conversion is assumed.')
                    row[field] = value.upper()
                else:
                    row[field] = value
            except (ValueError, OverflowError) as error:
                problems.append({'id': uuid4().hex, 'dataset_id': dataset['id'], 'file': dataset['name'], 'row': row_number,
                                 'field': field, 'cell': cell, 'value': value, 'message': str(error), 'severity': 'error'})
        if not problems and mapping.kind == 'orders' and row['received_qty'] > row['ordered_qty']:
            problems.append({'id': uuid4().hex, 'dataset_id': dataset['id'], 'file': dataset['name'], 'row': row_number,
                             'field': 'received_qty', 'cell': f"{get_column_letter(dataset['headers'].index(mapping.columns['received_qty'])+1)}{row_number}",
                             'value': str(row['received_qty']), 'message': 'Received exceeds ordered quantity; review the unit or order.', 'severity': 'error'})
        key = row.get('sku') if mapping.kind == 'inventory' else (row.get('order_id'), row.get('sku')) if mapping.kind == 'orders' else (row.get('sku'), row.get('date')) if mapping.kind == 'stock_history' else None
        if key is not None and not problems:
            if key in seen:
                problems.append({'id': uuid4().hex, 'dataset_id': dataset['id'], 'file': dataset['name'], 'row': row_number,
                                 'field': 'sku', 'cell': f"{get_column_letter(dataset['headers'].index(mapping.columns['sku'])+1)}{row_number}",
                                 'value': str(key), 'message': ('Repeated SKU/date history snapshot; later row quarantined. First occurrence retained.' if mapping.kind == 'stock_history' else 'Repeated inventory SKU or order line; later row quarantined. First occurrence retained.'), 'severity': 'error'})
            else:
                seen[key] = row_number
        if problems:
            quality.extend(problems)
        else:
            row['_source'] = {'dataset_id': dataset['id'], 'file': dataset['name'], 'sheet': dataset['sheet'], 'row': row_number,
                              'cells': {f: f"{get_column_letter(dataset['headers'].index(mapping.columns[f])+1)}{row_number}" for f in fields}}
            valid.append(row)
    return valid, quality
